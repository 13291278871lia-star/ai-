#!/usr/bin/env python3
"""Build an ordered, editable transcript review document with uncertain text in red."""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
VIDEO_ROOT = PROJECT / "assets/training/masters/videos"
TRANSCRIPT_ROOT = PROJECT / "assets/training/masters/subtitles"
OUTPUT = PROJECT / "assets/training/masters/高手分享字幕全文校对稿.docx"
UNCERTAIN = re.compile(r"【待確認：(.*?)】")


def natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def output_path(video: Path) -> Path:
    relative = video.relative_to(VIDEO_ROOT).with_suffix("")
    return Path(f"{TRANSCRIPT_ROOT / relative}.json")


def set_run_font(run, name: str = "STHeiti", size: float = 11, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_review_text(paragraph, text: str) -> int:
    cursor = 0
    count = 0
    for match in UNCERTAIN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        uncertain = match.group(1)
        run = paragraph.add_run(uncertain)
        set_run_font(run, bold=True)
        run.font.color.rgb = RGBColor(174, 34, 52)
        cursor = match.end()
        count += 1
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)
    return count


def add_page_number(paragraph) -> None:
    paragraph.alignment = 2
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9)


def main() -> None:
    videos = sorted(VIDEO_ROOT.rglob("*.mp4"), key=lambda path: natural_key(str(path.relative_to(VIDEO_ROOT))))
    records = []
    for video in videos:
        transcript = output_path(video)
        if not transcript.exists():
            continue
        try:
            data = json.loads(transcript.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        records.append((video, data))

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "STHeiti"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(16)
    title_run = title.add_run("高手分享字幕全文校对稿")
    set_run_font(title_run, size=24, bold=True)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(f"按合集 集数 分段及时间顺序整理  当前收录 {len(records)} 段视频  更新于 {datetime.now():%Y-%m-%d %H:%M}")
    set_run_font(run, size=10)
    run.font.color.rgb = RGBColor(100, 116, 139)
    intro = document.add_paragraph()
    add_review_text(intro, "本文档保留完整识别文本。红色文字表示语音识别把握较低，请对照对应视频时间修订；黑色文字也建议在正式使用前抽查。")

    total_uncertain = 0
    for index, (video, data) in enumerate(records):
        if index:
            document.add_page_break()
        relative = video.relative_to(VIDEO_ROOT)
        collection = relative.parts[0]
        heading = document.add_paragraph(style="Heading 1")
        heading_run = heading.add_run(collection)
        set_run_font(heading_run, size=11, bold=True)
        heading_run.font.color.rgb = RGBColor(148, 116, 52)
        video_title = document.add_paragraph(style="Heading 2")
        title_run = video_title.add_run(video.stem.replace(":", "/"))
        set_run_font(title_run, size=18, bold=True)
        title_run.font.color.rgb = RGBColor(0, 0, 0)

        for item in data.get("paragraphs", []):
            seconds = max(0, int(float(item.get("start", 0))))
            minute, second = divmod(seconds, 60)
            paragraph = document.add_paragraph()
            time_run = paragraph.add_run(f"[{minute:02d}:{second:02d}]  ")
            set_run_font(time_run, size=9, bold=True)
            time_run.font.color.rgb = RGBColor(148, 116, 52)
            total_uncertain += add_review_text(paragraph, str(item.get("text", "")))

    stats = document.paragraphs[1]
    stats.add_run(f"  待核对红字 {total_uncertain} 处")
    document.core_properties.title = "高手分享字幕全文校对稿"
    document.core_properties.subject = "繁体中文字幕人工校对"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT} with {len(records)} transcripts and {total_uncertain} uncertain spans")


if __name__ == "__main__":
    main()
