"""Build a source-grounded 11-page client deck inside the supplied GF template."""
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BURG = "681A1A"; GOLD = "8A7A49"; PALE = "F7F1E7"; INK = "1F1715"; GREY = "6E747C"; WHITE = "FFFFFF"

def _rgb(x): return RGBColor.from_string(x)

def _clear(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = _rgb(WHITE)

def _box(slide, x,y,w,h, fill=WHITE, line=None, radius=True):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=_rgb(fill)
    shp.line.color.rgb=_rgb(line or fill)
    return shp

def _text(slide, text, x,y,w,h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial Unicode MS", valign=MSO_ANCHOR.MIDDLE):
    shp=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=shp.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=valign; tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.text=str(text); p.alignment=align
    for r in p.runs: r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=_rgb(color)
    return shp

def _title(slide, kicker, title, subtitle=""):
    _text(slide,kicker.upper(),.55,.28,5,.25,10,BURG,True)
    _text(slide,title,.55,.58,12.1,.62,26,INK,True)
    if subtitle: _text(slide,subtitle,.58,1.20,12,.35,11,GREY)
    _box(slide,.55,1.58,1.0,.05,GOLD,radius=False)

def _foot(slide, source, fx=None):
    fxtext=f"｜人民币按 1 美元 = {fx:g} 人民币展示" if fx else ""
    _text(slide,f"资料来源：{source}{fxtext}｜数值仅供参考，详见计划书。",.55,7.12,12.2,.22,7,GREY)

def _table(slide, data, x,y,w,h, widths=None, font=12):
    rows=len(data); cols=len(data[0]); shape=slide.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(h)); tbl=shape.table
    if widths:
        for i,cw in enumerate(widths): tbl.columns[i].width=Inches(cw)
    for ri,row in enumerate(data):
        for ci,val in enumerate(row):
            cell=tbl.cell(ri,ci); cell.text=str(val); cell.margin_left=cell.margin_right=Inches(.05); cell.margin_top=cell.margin_bottom=Inches(.03)
            cell.fill.solid(); cell.fill.fore_color.rgb=_rgb(GOLD if ri==0 else ("FBF8F2" if ri%2 else WHITE))
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment=PP_ALIGN.CENTER
                for r in p.runs: r.font.name="Arial Unicode MS"; r.font.size=Pt(font); r.font.bold=(ri==0); r.font.color.rgb=_rgb(WHITE if ri==0 else INK)
    return shape

def _metric(slide,label,value,x,y,w=2.6):
    _box(slide,x,y,w,.9,PALE,GOLD)
    _text(slide,label,x+.15,y+.12,w-.3,.22,10,GREY)
    _text(slide,value,x+.15,y+.37,w-.3,.38,19,BURG,True)

def _bullet(slide, title, body, x,y,w=5.7,h=.82):
    _box(slide,x,y,w,h,"FBF8F2","E7DCC8")
    _text(slide,title,x+.18,y+.10,w-.36,.23,13,BURG,True)
    _text(slide,body,x+.18,y+.35,w-.36,h-.42,10,INK)

def _move_slide(prs, old, new):
    ids=prs.slides._sldIdLst; item=ids[old]; ids.remove(item); ids.insert(new,item)

def build_template_ppt(profile, rows, template_path):
    prs=Presentation(str(template_path)); source=str(profile.get("source") or "上传计划书")[:110]
    fx=float(profile.get("fx") or 6.8); premium=float(profile.get("premium") or 0); levy=float(profile.get("levy") or 0); years=int(profile.get("years") or 1)
    age=int(profile.get("age") or 0); product=str(profile.get("product") or "保险计划"); name=str(profile.get("insured") or "受保人"); sum_assured=float(profile.get("sumAssured") or 0)
    by={int(r["year"]):r for r in rows}; milestones=[y for y in (20,30,40,50,60) if y in by] or [int(r["year"]) for r in rows[-5:]]
    total_outlay=premium*years+levy
    asset=Path(template_path).parent.parent/"generated"
    img1=asset/"wealth-legacy-family.png"; img2=asset/"long-term-planning.png"
    while len(prs.slides)<11: prs.slides.add_slide(prs.slide_layouts[6])

    # 1 — Overview
    s=prs.slides[0]; _clear(s)
    _text(s,"香港友邦保险",.45,.35,3.0,.45,26,INK,True); _text(s,product,3.0,.35,9.7,.5,25,GOLD,True)
    _text(s,"终身保障与长期财富传承规划",.5,.95,7,.28,12,GREY)
    _text(s,"[ 计划重点 ]",.5,1.65,4.5,.45,23,INK,True); _text(s,"[ 预期现金价值 ]",5.4,1.65,7.2,.45,23,INK,True)
    _metric(s,"整付保费",f"USD {premium:,.0f}",.65,2.35,3.9); _metric(s,"首期总支出",f"USD {total_outlay:,.2f}",.65,3.45,3.9)
    _metric(s,"基本保额",f"USD {sum_assured:,.0f}" if sum_assured else "计划书提取中",.65,4.55,3.9)
    _text(s,f"{name}｜{age}岁｜美元保单｜{'整付保费' if years==1 else str(years)+'年供款'}",.7,5.65,4.0,.65,13,INK,True)
    data=[["保单年末","美元","人民币","年化 IRR"]]
    for y in milestones:
        r=by[y]; data.append([f"{y}年 / {age+y}岁",f"{r['usd']:,.0f}",f"{r['usd']*fx:,.0f}",f"{r['rate']*100:.2f}%"])
    _table(s,data,5.35,2.3,7.45,3.75,[1.55,1.75,2.25,1.9],12)
    _text(s,"IRR 以整付保费及保费征费为初始现金流，各年为独立退保情景。",5.45,6.25,7.2,.3,9,GREY)
    _foot(s,source,fx)

    # 2 — Editable IRR chart
    s=prs.slides[1]; _clear(s); _title(s,"RETURN PROFILE","不同年份的 IRR 走势",f"{name} · {age}岁投保 · 初始支出 USD {total_outlay:,.2f}")
    chart_rows=[by[y] for y in milestones]
    cd=CategoryChartData(); cd.categories=[f"{int(r['year'])}年" for r in chart_rows]; cd.add_series("年化 IRR",[float(r["rate"])*100 for r in chart_rows])
    chart=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(.8),Inches(1.9),Inches(11.8),Inches(4.55),cd).chart
    chart.has_legend=False; chart.value_axis.tick_labels.number_format='0.0"%"'; chart.value_axis.minimum_scale=0; chart.value_axis.has_major_gridlines=True; chart.category_axis.tick_labels.font.size=Pt(8); chart.value_axis.tick_labels.font.size=Pt(9)
    series=chart.series[0]; series.format.line.color.rgb=_rgb(BURG); series.format.line.width=Pt(2.5); series.marker.size=7
    _text(s,"退保价值包含非保证终期分红，IRR 并非保证回报。",.82,6.55,11.5,.3,10,BURG,True); _foot(s,source)

    # 3 — milestone table
    s=prs.slides[2]; _clear(s); _title(s,"VALUE MILESTONES","预期退保价值里程碑","保证价值与非保证终期分红分开呈现")
    data=[["年末 / 年龄","保证金额","非保证分红","预期总额"]]
    for y in milestones:
        r=by[y]; data.append([f"{y}年 / {age+y}岁",f"USD {float(r.get('guaranteed') or 0):,.0f}",f"USD {float(r.get('bonus') or 0):,.0f}",f"USD {r['usd']:,.0f}"])
    _table(s,data,.7,1.9,11.9,4.35,[2.1,2.8,3.2,3.2],14); _foot(s,source)

    # 4 — legacy visual
    s=prs.slides[3]; _clear(s)
    if img1.exists(): s.shapes.add_picture(str(img1),Inches(6.1),Inches(0),width=Inches(7.23),height=Inches(7.5))
    _box(s,0,0,6.55,7.5,WHITE,radius=False); _title(s,"LEGACY","跨代传承与身故保障")
    _bullet(s,"首 10 个保单年度","基本保额的 100% 加终期分红的面值。",.65,2.0,5.1,.9)
    _bullet(s,"第 10 年后","取“已缴整付保费”与“基本保额 × 适用百分比”之较高者，再加终期分红。",.65,3.1,5.1,1.05)
    _bullet(s,"适用百分比","其后按计划书规则逐年下降，最终至基本保额的 50%。",.65,4.38,5.1,.9)
    _text(s,"实际身故赔偿以保单条款及当时保单状态为准。",.7,5.65,5,.45,10,GREY); _foot(s,source)

    # 5 — continuity
    s=prs.slides[4]; _clear(s); _title(s,"POLICY CONTINUITY","保单延续与传承安排","计划书列明的保单管理选择")
    items=[("第二持有人","让指定人士在原持有人身故后承接保单。"),("身故赔偿支付选择","可按计划条款安排支付方式。"),("受益人灵活选择","为不同家庭阶段预留调整空间。"),("保单分拆选择","符合条件时可按条款分拆保单。")]
    for i,(a,b) in enumerate(items): _bullet(s,a,b,.75+(i%2)*6.05,2.0+(i//2)*1.55,5.55,1.15)
    _foot(s,source)

    # 6 — flexibility and visual
    s=prs.slides[5]; _clear(s)
    if img2.exists(): s.shapes.add_picture(str(img2),Inches(0),Inches(0),width=Inches(6.4),height=Inches(7.5))
    _box(s,5.85,0,7.48,7.5,WHITE,radius=False); _title(s,"FLEXIBILITY","长期规划的弹性工具")
    for i,(a,b) in enumerate([("健康障碍选择","在符合条款的情况下提供额外安排。"),("预先核保选择","受保人 60 岁或以下时，可按计划书所述条件使用。"),("保单贷款","可按当时保单价值及公司条款申请。"),("保单逆按揭计划","是否合资格及获批由相关机构决定。")]): _bullet(s,a,b,6.2,1.8+i*1.18,6.35,.92)
    _foot(s,source)

    # 7 — assumptions
    s=prs.slides[6]; _clear(s); _title(s,"ASSUMPTIONS","本次演示的计算口径","只采用这份计划书的无提款基本说明")
    facts=[("现金流",f"t=0 支出 USD {total_outlay:,.2f}；指定年末收回退保总额。"),("提款", "计划书没有提供中途提款方案，本简报未自行设计提款。"),("分红","终期分红属非保证，可能高于或低于演示数值。"),("汇率",f"人民币只按 {fx:g} 展示，不属于保单回报。"),("IRR","每个年期独立计算，不能把不同年期的退保价值相加。")]
    for i,(a,b) in enumerate(facts): _bullet(s,a,b,.8,1.85+i*.95,11.7,.72)
    _foot(s,source)

    # 8–10 — unavailable template-specific scenarios, transparently handled
    pages=[
      ("WITHDRAWAL SCENARIO","提款情景","本计划书未提供提款金额、提款年龄或提款后价值，因此系统不自动填充。","如需要此页，请上传保险公司列明提款安排的同一产品计划书。"),
      ("PROMOTION","推广与预缴优惠","本计划书没有现行推广期、回赠或预缴利率资料，因此系统不沿用模板旧数字。","推广资料需与产品、币种、保费及推广期逐项配对。"),
      ("COMPARISON","方案比较","目前只有一个已核对方案，不能生成不选择预缴／预缴 1 年／预缴 4 年比较。","上传对应推广资料后，平台可按同一汇率与口径生成比较。")]
    for idx,(kick,ttl,main,sub) in enumerate(pages,7):
        s=prs.slides[idx]; _clear(s); _title(s,kick,ttl,"资料完整性检查")
        _box(s,1.0,2.05,11.3,2.55,PALE,"D7C59D"); _text(s,"未从计划书提取",1.5,2.45,10.3,.55,25,BURG,True,PP_ALIGN.CENTER); _text(s,main,1.6,3.15,10.1,.6,14,INK,False,PP_ALIGN.CENTER); _text(s,sub,1.6,4.0,10.1,.35,10,GREY,False,PP_ALIGN.CENTER); _foot(s,source)

    # 11 — source and cautions
    s=prs.slides[10]; _clear(s); _title(s,"SOURCE CHECK","资料来源与重要提示","自动生成前已执行产品、币种与口径隔离")
    checks=[("产品",product),("受保人",f"{name}，投保年龄 {age} 岁"),("保费",f"{'整付' if years==1 else str(years)+'年供款'} USD {premium:,.2f}；征费 USD {levy:,.2f}"),("提款假设","没有提款"),("非保证元素","终期分红及相关预期价值并非保证"),("风险","早期退保可能导致重大损失；人民币展示受汇率波动影响")]
    _table(s,[["核对项目","本次生成采用的资料"]]+checks,.8,1.85,11.7,4.65,[2.6,9.1],12); _foot(s,source)

    bio=BytesIO(); prs.save(bio); return bio.getvalue()
