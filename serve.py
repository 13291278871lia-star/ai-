#!/usr/bin/env python3
"""Run the local MVP at http://127.0.0.1:4175."""
import json
import os
import re
import unicodedata
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse


def build_template_ppt(profile, rows, template_path):
    """Open the user's GF template, fill slide-1 table/text and add a native
    editable line-chart slide. Returns the saved .pptx as bytes.

    pptxgenjs (used client-side) can only *generate* a fresh deck; it cannot
    open an existing .pptx. Filling the real 11-page template therefore has to
    happen server-side with python-pptx, which is why this lives here.
    """
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN

    fx = float(profile.get("fx", 6.8) or 6.8)
    premium = float(profile.get("premium", 0) or 0)
    years = int(profile.get("years", 1) or 1)
    product = str(profile.get("product", "") or "")
    source = str(profile.get("source", "") or "")
    age = profile.get("age", "")
    pay_term = "整付保費" if years == 1 else f"{years}年"

    prs = Presentation(str(template_path))
    slide0 = prs.slides[0]

    def set_textbox(shape, lines):
        tf = shape.text_frame
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        first = tf.paragraphs[0]
        for i, line in enumerate(lines):
            if i == 0:
                first.text = line
            else:
                tf.add_paragraph().text = line

    by_year = {int(r["year"]): r for r in rows if "year" in r}

    # --- fill slide-1 text placeholders (matched by shape name) ---
    for shape in slide0.shapes:
        if not shape.has_text_frame:
            continue
        nm = shape.name
        if nm == "文字方塊 4":  # main title
            set_textbox(shape, [f"香港友邦保险 - {product}"])
        elif nm == "文字方塊 6":  # 目标 / 优点
            set_textbox(shape, [
                "目标：终身分红保障，财富传承与灵活现金提取并重",
                "优点：整付保费锁定价值，享终期分红非保证收益",
            ])
        elif nm == "文字方塊 8":  # 供款额 / 存款期 / 总供款
            set_textbox(shape, [
                f"供款额：{premium:,.0f} 美元 ({premium * fx:,.0f} 人民币)",
                f"存款期：{pay_term}    总供款：{premium * years:,.0f} 美元 ({premium * years * fx:,.0f} 人民币)",
            ])

    # --- fill slide-1 summary table (表格 1): 20/30/40/50年后 ---
    for shape in slide0.shapes:
        if shape.has_table and shape.name == "表格 1":
            tbl = shape.table
            for ri, yr in enumerate((20, 30, 40, 50), start=1):
                r = by_year.get(yr)
                if not r:
                    continue
                usd = float(r.get("usd", 0) or 0)
                rmb = float(r.get("rmb", usd * fx) or usd * fx)
                rate = float(r.get("rate", 0) or 0)
                tbl.cell(ri, 0).text = f"{yr}年后"
                tbl.cell(ri, 1).text = f"{usd:,.0f}"
                tbl.cell(ri, 2).text = f"{rmb:,.0f}"
                tbl.cell(ri, 3).text = f"{rate * 100:.2f}%"
                for c in range(4):
                    for p in tbl.cell(ri, c).text_frame.paragraphs:
                        p.alignment = PP_ALIGN.CENTER
                        for run in p.runs:
                            run.font.size = Pt(14)
                            run.font.bold = True
            break

    # --- add a native, editable line-chart slide (value growth) ---
    sorted_rows = sorted(rows, key=lambda r: r["year"])
    chart_data = CategoryChartData()
    chart_data.categories = [f"第{int(r['year'])}年" for r in sorted_rows]
    has_split = any(r.get("guaranteed") is not None for r in sorted_rows)
    if has_split:
        chart_data.add_series("保证金额", [float(r.get("guaranteed") or 0) for r in sorted_rows])
        chart_data.add_series("非保证分红", [float(r.get("bonus") or 0) for r in sorted_rows])
    chart_data.add_series("退保总额", [float(r.get("usd", 0) or 0) for r in sorted_rows])

    blank = None
    for layout in prs.slide_layouts:
        if "blank" in layout.name.lower() or "空白" in layout.name:
            blank = layout
            break
    chart_slide = prs.slides.add_slide(blank or prs.slide_layouts[6])

    title = f"{product} — 预期退保价值增长（{premium:,.0f}美元{pay_term}）"
    tx = chart_slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.8))
    tx.text_frame.word_wrap = True
    tp = tx.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(22)
    tp.font.bold = True

    chart_shape = chart_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(0.4), Inches(1.1), Inches(12.5), Inches(5.8), chart_data
    )
    chart = chart_shape.chart
    chart.has_title = True
    chart.chart_title.text_frame.paragraphs[0].text = "预期退保价值增长曲线（美元）"
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(16)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.plots[0].has_data_labels = False
    chart.value_axis.tick_labels.number_format = "#,##0"
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.paragraphs[0].text = "金额 (美元)"
    chart.value_axis.axis_title.text_frame.paragraphs[0].font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(8)

    foot = chart_slide.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(12.5), Inches(0.4))
    fp = foot.text_frame.paragraphs[0]
    fp.text = f"*以美元兑人民币{fx}计算*  以上数据只供参考, 详情请参阅建议书。来源：{source}"
    fp.font.size = Pt(9)
    fp.font.italic = True

    bio = BytesIO()
    prs.save(bio)
    return bio.getvalue()


class AppHandler(SimpleHTTPRequestHandler):
    extensions_map = SimpleHTTPRequestHandler.extensions_map | {
        ".js": "application/javascript",
        ".mjs": "application/javascript",
    }

    def end_headers(self):
        # The MVP is edited frequently. Avoid showing colleagues an older cached UI.
        self.send_header("Cache-Control", "no-store, max-age=0")
        super().end_headers()

    def do_GET(self):
        request = urlparse(self.path)
        if request.path == "/api/library-files":
            self.send_library_files(parse_qs(request.query))
            return
        if request.path == "/api/product-content":
            self.send_product_content(parse_qs(request.query))
            return
        super().do_GET()

    def do_POST(self):
        request = urlparse(self.path)
        if request.path == "/api/ppt-from-template":
            self.send_template_ppt()
            return
        self.send_error(404, "Not found")

    def send_template_ppt(self):
        """Fill the user's GF template (server-side, python-pptx) and return a
        full editable .pptx with the summary table filled and a native
        line-chart slide added. The browser cannot do this with pptxgenjs
        because that library can only build a deck from scratch, not open an
        existing template."""
        length = int(self.headers.get("Content-Length", 0) or 0)
        if length <= 0 or length > 2_000_000:
            self.send_error(400, "Invalid request size")
            return
        raw = self.rfile.read(length)
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError:
            self.send_error(400, "Invalid JSON")
            return
        profile = payload.get("profile") if isinstance(payload, dict) else None
        rows = payload.get("rows") if isinstance(payload, dict) else None
        if not isinstance(profile, dict) or not isinstance(rows, list) or not rows:
            self.send_error(400, "Missing profile or rows")
            return

        template_path = Path("assets/ppt-reference/gf-template.pptx")
        site_root = Path.cwd().resolve()
        template_path = (site_root / template_path).resolve()
        try:
            template_path.relative_to(site_root)
        except ValueError:
            self.send_error(400, "Invalid template path")
            return
        if not template_path.is_file():
            self.send_error(404, "GF template not found at assets/ppt-reference/gf-template.pptx")
            return

        try:
            import pptx  # noqa: F401
        except ImportError:
            self.send_error(500, "python-pptx not installed. Run: pip install python-pptx")
            return
        try:
            data = build_template_ppt(profile, rows, template_path)
        except Exception as exc:  # surface the real cause to the browser
            self.send_error(500, f"Template build failed: {exc}")
            return

        fname = (str(profile.get("product") or "overview").replace("/", "_"))[:60]
        disp_name = f"{fname}_完整模版.pptx"
        # HTTP headers are latin-1; provide an ASCII fallback plus a UTF-8
        # filename* (RFC 5987) so browsers get the proper Chinese name.
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.send_header("Content-Disposition", f'attachment; filename="policy_overview_template.pptx"; filename*=UTF-8\'\'{quote(disp_name)}')
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def send_library_files(self, query):
        library = query.get("library", [""])[0]
        folder = query.get("folder", [""])[0]
        roots = {
            "tax": Path("assets/library/01_友扣稅"),
            "health": Path("assets/library/02_AIA健康系列"),
        }
        base = roots.get(library)
        if not base or not folder:
            self.send_error(400, "Invalid library request")
            return

        site_root = Path.cwd().resolve()
        base_path = (site_root / base).resolve()
        target = (base_path / folder).resolve()
        try:
            target.relative_to(base_path)
        except ValueError:
            self.send_error(400, "Invalid folder")
            return
        if not target.is_dir():
            self.send_error(404, "Product folder not found")
            return

        files = [
            {"name": item.name, "path": item.relative_to(site_root).as_posix()}
            for item in sorted(target.iterdir(), key=lambda item: item.name.casefold())
            if item.is_file() and item.suffix.lower() == ".pdf"
        ]
        body = json.dumps({"files": files}, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def send_product_content(self, query):
        library = query.get("library", [""])[0]
        folder = query.get("folder", [""])[0]
        try:
            fields = json.loads(query.get("fields", ["[]"])[0])
        except json.JSONDecodeError:
            self.send_error(400, "Invalid field request")
            return
        if not isinstance(fields, list) or len(fields) > 32:
            self.send_error(400, "Invalid field request")
            return

        site_root = Path.cwd().resolve()
        library_root = (site_root / "assets" / "library" / library).resolve()
        target = (library_root / folder).resolve()
        try:
            target.relative_to(library_root)
        except ValueError:
            self.send_error(400, "Invalid product request")
            return
        if not library_root.is_dir() or not target.is_dir():
            self.send_error(404, "Product folder not found")
            return

        clean_fields = []
        for field in fields:
            if not isinstance(field, dict):
                continue
            key, label, aliases = field.get("key"), field.get("label"), field.get("aliases", [])
            if not isinstance(key, str) or not isinstance(label, str):
                continue
            if not isinstance(aliases, list):
                aliases = []
            clean_keywords = [word for word in aliases[:16] if isinstance(word, str) and 0 < len(word) <= 30]
            if clean_keywords:
                clean_fields.append({"key": key[:60], "label": label[:80], "keywords": clean_keywords})

        try:
            from pypdf import PdfReader
        except ImportError:
            self.send_error(500, "PDF reader unavailable")
            return

        def normalized(value):
            value = unicodedata.normalize("NFKC", value).translate(str.maketrans({
                "終": "终", "賠": "赔", "額": "额", "醫": "医", "療": "疗", "墊": "垫",
                "費": "费", "類": "类", "別": "别", "術": "术", "護": "护", "計": "计", "劃": "划",
            }))
            return re.sub(r"[\s\W_]+", "", value)

        extracted = {field["key"]: [] for field in clean_fields}
        source_files = []
        for item in sorted(target.iterdir(), key=lambda entry: entry.name.casefold()):
            if not item.is_file() or item.suffix.lower() != ".pdf":
                continue
            try:
                reader = PdfReader(str(item))
                lines = []
                for page in reader.pages:
                    text = page.extract_text() or ""
                    lines.extend(line.strip() for line in text.splitlines() if line.strip())
                source_files.append(item.name)
            except Exception:
                continue
            for field in clean_fields:
                normalized_keywords = [normalized(word) for word in field["keywords"]]
                for index, line in enumerate(lines):
                    if any(word in normalized(line) for word in normalized_keywords):
                        evidence = " ".join(lines[max(0, index - 1):min(len(lines), index + 2)])[:360]
                        value = " ".join(lines[index:min(len(lines), index + 3)])[:220]
                        extracted[field["key"]].append({"value": value, "evidence": evidence, "source": item.name, "found": True})
                        break

        result = {
            "productName": folder,
            "textRead": bool(source_files),
            "sourceFiles": source_files,
            "fields": {
                field["key"]: {"label": field["label"], "results": extracted[field["key"]] or [{"value": "資料暫未提供", "evidence": "", "source": "", "found": False}]}
                for field in clean_fields
            },
        }
        body = json.dumps(result, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "4175"))
    server = ThreadingHTTPServer(("127.0.0.1", port), AppHandler)
    print(f"AI 工作台 MVP 已启动：http://127.0.0.1:{port}")
    server.serve_forever()
